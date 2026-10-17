import csv
import hashlib
import io
import json
import zipfile
from decimal import Decimal, InvalidOperation
import pandas as pd
from django.core.exceptions import ValidationError
from django.db import transaction
from .models import DataImport, Transaction
from .access import audit

MAX_SIZE = 5 * 1024 * 1024
FIELDS = ["date", "description", "amount", "kind", "currency", "status", "category", "reference", "due_date"]


def preview_import(business, upload, mapping=None):
    if upload.size > MAX_SIZE:
        raise ValidationError("File limit is 5 MB.")
    raw = upload.read(MAX_SIZE + 1)
    if len(raw) > MAX_SIZE:
        raise ValidationError("File limit is 5 MB.")
    digest = hashlib.sha256(raw + json.dumps(mapping or {}, sort_keys=True).encode()).hexdigest()
    existing = DataImport.objects.filter(business=business, digest=digest).first()
    if existing:
        return existing
    try:
        if upload.name.lower().endswith(".csv"):
            if b"\x00" in raw:
                raise ValueError("Binary content is not CSV.")
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            source = []
            for i, row in enumerate(reader):
                if i >= 5000:
                    raise ValueError("Maximum 5,000 rows per import.")
                source.append(row)
        elif upload.name.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 25 * 1024 * 1024 or len(archive.infolist()) > 1000:
                    raise ValueError("Workbook expands beyond safe limits.")
                if any("vbaProject" in x.filename or "externalLinks/" in x.filename for x in archive.infolist()):
                    raise ValueError("External links and macros are not allowed.")
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheet = book.active
            values = sheet.iter_rows()
            headers = [str(c.value or "").strip() for c in next(values)]
            source = []
            for i, cells in enumerate(values):
                if i >= 5000:
                    raise ValueError("Maximum 5,000 rows per import.")
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Formula cells are not accepted.")
                source.append(dict(zip(headers, [c.value for c in cells])))
            book.close()
        else:
            raise ValueError("Upload a UTF-8 CSV or XLSX workbook.")
    except (ValueError, UnicodeError, zipfile.BadZipFile, StopIteration, KeyError) as exc:
        raise ValidationError(str(exc)) from exc
    rows, errors, seen = [], [], set()
    for index, original in enumerate(source, 2):
        row = {field: original.get((mapping or {}).get(field) or field, "") for field in FIELDS}
        try:
            day = pd.to_datetime(str(row["date"]), format="ISO8601", errors="raise").date()
            amount = Decimal(str(row["amount"]))
            if not amount.is_finite() or amount <= 0 or amount >= Decimal("10000000000000000") or amount.as_tuple().exponent < -2:
                raise ValueError("Amount must be positive with at most two decimal places.")
            kind = str(row["kind"]).lower().strip()
            if kind not in dict(Transaction.TYPES):
                raise ValueError("Unknown transaction type.")
            currency = str(row["currency"]).strip().upper()
            if currency != business.currency:
                raise ValueError("Currency must match the workspace.")
            status = str(row["status"]).lower().strip()
            if status not in ("paid", "pending"):
                raise ValueError("Explicit paid or pending status is required.")
            if day < business.opening_date:
                raise ValueError("Date precedes opening balance date.")
            due = pd.to_datetime(str(row["due_date"]), format="ISO8601").date().isoformat() if row["due_date"] else None
            if status == "pending" and (not due or kind not in ("income", "expense", "receivable", "payable")):
                raise ValueError("Pending rows require a valid type and due date.")
            description = str(row["description"]).strip()
            if not description or len(description) > 240:
                raise ValueError("Description is required (maximum 240 characters).")
            category, reference = str(row["category"] or "General"), str(row["reference"] or "")
            if len(category) > 80 or len(reference) > 100:
                raise ValueError("Category or reference is too long.")
            normalized = {
                "date": day.isoformat(),
                "description": description,
                "amount": str(amount.quantize(Decimal("0.01"))),
                "kind": kind,
                "currency": currency,
                "status": status,
                "category": category,
                "reference": reference,
                "due_date": due,
            }
            fingerprint = hashlib.sha256(json.dumps(normalized, sort_keys=True).encode()).hexdigest()
            if fingerprint in seen or Transaction.objects.filter(business=business, fingerprint=fingerprint).exists():
                raise ValueError("Duplicate row detected; no rows will be imported.")
            seen.add(fingerprint)
            normalized["fingerprint"] = fingerprint
            rows.append(normalized)
        except (ValueError, TypeError, InvalidOperation, OverflowError) as exc:
            errors.append({"row": index, "error": str(exc)[:300]})
    if not source:
        errors.append({"row": 1, "error": "No data rows found."})
    return DataImport.objects.create(business=business, digest=digest, rows=rows, errors=errors)


@transaction.atomic
def confirm_import(business, user, import_id):
    item = DataImport.objects.select_for_update().get(pk=import_id, business=business)
    if item.errors:
        raise ValidationError("Resolve all row errors and upload a corrected file.")
    if item.confirmed:
        return 0
    account = business.accounts.first()
    count = 0
    for row in item.rows:
        _, created = Transaction.objects.get_or_create(
            business=business,
            fingerprint=row["fingerprint"],
            defaults={**{k: v for k, v in row.items() if k != "fingerprint"}, "account": account},
        )
        count += int(created)
    item.confirmed = True
    item.save(update_fields=["confirmed"])
    audit(user, business, "import.confirmed", item.pk)
    return count


def safe_cell(value):
    text = str(value if value is not None else "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text
