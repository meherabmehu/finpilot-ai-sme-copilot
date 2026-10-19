import csv
import io
from decimal import Decimal
from django.http import HttpResponse, HttpResponseBadRequest
from django.utils import timezone
from .engine import summary
from .imports import safe_cell
from .access import audit, authorized_business
from .report_sections import agentic_sections, render_pdf, amount


def export_report(user, business, format, start, end, report_type="summary"):
    business = authorized_business(user, business.pk)
    if report_type in ('transactions', 'receivables', 'payables') and business.transactions.filter(archived=False, date__gte=start, date__lte=end).count() > 20000:
        return HttpResponseBadRequest('This export exceeds 20,000 source records. Narrow the reporting period; no partial file was generated.')
    if report_type == 'scenarios' and business.scenarios.count() > 500:
        return HttpResponseBadRequest('More than 500 scenario snapshots require a reviewed offline export; no partial report was generated.')
    values = summary(business, start, end)
    rows = [
        ["FinPilot financial report", business.name],
        ["Period", f"{start} to {end}"],
        ["Currency", business.currency],
        ["Generated UTC", timezone.now().isoformat()],
        ["Source", "Synthetic demonstration records" if business.is_demo else "User-entered records; completeness unverified"],
        ["Basis", "Recorded operating summary; not audited net profit. Cash balance is as of period end."],
    ]
    sections = None
    from .intelligence.financial_state import build_state
    source = build_state(business)
    rows += [['Report type', report_type], ['Current source fingerprint (not archived decision evidence)', source['fingerprint']],
             ['Completeness', str(source['complete'])], ['Currency warning', values['mixed_currency_warning'] or 'No currency conversion applied.']]
    metadata = list(rows)
    if report_type in ('survival', 'scenarios', 'roadmap'):
        from .survival_report_sections import survival_sections, scenario_sections, roadmap_sections
        sections = {'survival': survival_sections, 'scenarios': scenario_sections, 'roadmap': roadmap_sections}[report_type](business)
        for section in sections:
            rows.append([section['title']])
            rows.append(section['headers'])
            rows.extend(section['rows'])
    elif report_type in ('decisions', 'investigations', 'issues'):
        rows.append(['Basis', 'SIMULATED alternatives are not observed results. No causal improvement or external execution is inferred.'])
        sections = agentic_sections(business, start, end, report_type)
        for section in sections:
            rows.append([section['title']])
            rows.append(section['headers'])
            rows.extend(section['rows'])
    elif report_type in ('receivables', 'payables'):
        kind = 'receivable' if report_type == 'receivables' else 'payable'
        rows.append(['Basis', 'Pending recorded invoices, filtered by record date. Unknown due dates remain unknown; no collection or payment is guaranteed.'])
        rows.append(['Record date', 'Description', 'Amount', 'Currency', 'Due date', 'Reference', 'Timing status', 'Days overdue', 'Aging bucket'])
        rows.extend([
            [r.date, r.description, r.amount, r.currency, r.due_date or 'Unknown', r.reference,
             'Unknown due date' if not r.due_date else 'Overdue' if r.due_date < timezone.localdate() else 'Upcoming / due today',
             max((timezone.localdate()-r.due_date).days, 0) if r.due_date else 'Unknown',
             'Unknown due date' if not r.due_date else 'Not overdue' if r.due_date >= timezone.localdate() else
             '1-30 days' if (timezone.localdate()-r.due_date).days <= 30 else '31-60 days' if (timezone.localdate()-r.due_date).days <= 60 else '61-90 days' if (timezone.localdate()-r.due_date).days <= 90 else '90+ days']
            for r in business.transactions.filter(archived=False, kind=kind, status='pending', date__gte=start, date__lte=end)
        ])
    elif report_type == "transactions":
        rows.append(["Date", "Description", "Type", "Amount", "Currency", "Status", "Category", "Reference"])
        rows.extend(
            [
                [r.date, r.description, r.kind, r.amount, r.currency, r.status, r.category, r.reference]
                for r in business.transactions.filter(archived=False, date__gte=start, date__lte=end)
            ]
        )
    elif report_type == "forecast":
        from .intelligence.forecasting import forecast_with_evaluation
        result = forecast_with_evaluation(business)
        rows.append(["Method", result.get("limitations", result.get("reason"))])
        rows.append(["Date", "Projected balance", "Expected inflow", "Expected outflow"])
        rows.extend([[p["date"], amount(p["balance"]), amount(p["inflow"]), amount(p["outflow"])] for p in result.get("points", [])])
    elif report_type == "financing":
        rows.append(["Saved UTC", "Principal", "Annual rate", "Months", "Monthly payment", "Total cost", "Fees"])
        rows.extend(
            [
                [
                    s.created_at.isoformat(),
                    amount(s.assumptions["amount"]),
                    s.assumptions["annual_rate"] + '%',
                    int(s.assumptions["months"]),
                    amount(s.result["payment"]),
                    amount(s.result["total_cost"]),
                    amount(s.result["fees"]),
                ]
                for s in business.loan_comparisons.all()
            ]
        )
        rows.append(
            ["Limitations", "User-entered illustrations; not loan offers or approvals. Reducing-balance monthly repayments; upfront fees."]
        )
    elif report_type == "quality":
        rows.extend(
            [
                ["Record count", values["count"]],
                ["Cash records user-confirmed complete", business.cash_records_complete],
                ["Opening balance known", business.opening_balance is not None],
                ["Accounting completeness", "Not independently verified; inventory, taxes and accrual completeness are unknown."],
            ]
        )
    else:
        rows.extend([[k.replace("_", " ").title(), v if v is not None else "Unknown"] for k, v in values.items()])
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerows([[safe_cell(x) for x in row] for row in rows])
        response = HttpResponse("\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8")
    elif format == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Financial report"
        from openpyxl.styles import Font, PatternFill
        for row in rows:
            sheet.append([x if isinstance(x, (Decimal, int, float, bool)) else safe_cell(x) for x in row])
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, (Decimal, float)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00)'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='173D7A')
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 60
        output = io.BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    elif format == "pdf":
        if sections is None:
            # Preserve all existing exports but wrap their rows in readable paginated tables.
            width = max(len(row) for row in rows[6:]) if len(rows) > 6 else 2
            sections = [{'title': report_type.replace('_', ' ').title(),
                         'headers': ['Measure / record'] + ['Value']*(width-1), 'rows': rows[6:]}]
        response = HttpResponse(render_pdf(business, start, end, metadata, sections), content_type='application/pdf')
    else:
        raise ValueError("Unsupported export format")
    response["Content-Disposition"] = f'attachment; filename="finpilot-{report_type}.{format}"'
    audit(user, business, "report.exported", report_type + "." + format)
    return response
